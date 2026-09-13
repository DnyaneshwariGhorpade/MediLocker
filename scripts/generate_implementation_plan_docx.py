"""Render IMPLEMENTATION_PLAN.md into a styled Word document.

Reuses the table/callout styling helpers from generate_docx.py so the plan
matches the look of the existing MediLocker specification documents.

Usage:  python scripts/generate_implementation_plan_docx.py
"""

import os
import re
import sys

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate_docx import (  # noqa: E402
    add_styled_table,
    set_cell_background,
    set_cell_margins,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = os.path.join(ROOT, "IMPLEMENTATION_PLAN.md")
OUTPUTS = [
    os.path.join(ROOT, "MediLocker_Implementation_Plan.docx"),
    os.path.join(ROOT, "docs", "MediLocker_Implementation_Plan.docx"),
]

NAVY = RGBColor(30, 58, 138)
BLUE = RGBColor(37, 99, 235)
SLATE = RGBColor(30, 41, 59)
GREY = RGBColor(100, 116, 139)

HEADING_SIZES = {1: 16, 2: 14, 3: 12}


# ---------------------------------------------------------------- inline text

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10, color=SLATE):
    """Writes text into a paragraph, honouring **bold** and `code` spans."""
    for token in TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
            run.font.name = "Segoe UI"
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(size - 1)
            run.font.color.rgb = RGBColor(190, 24, 93)
            continue
        else:
            run = paragraph.add_run(token)
            run.font.name = "Segoe UI"
        run.font.size = Pt(size)
        run.font.color.rgb = color


# ------------------------------------------------------------------- markdown


def strip_cell(text):
    return text.replace("**", "").replace("`", "").strip()


def parse_table(lines, index):
    """Consumes a pipe table starting at `index`; returns (headers, rows, next)."""
    header = [strip_cell(c) for c in lines[index].strip().strip("|").split("|")]
    index += 2  # skip the separator row
    rows = []
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        cells = [strip_cell(c) for c in lines[index].strip().strip("|").split("|")]
        cells = (cells + [""] * len(header))[: len(header)]
        rows.append(cells)
        index += 1
    return header, rows, index


def column_widths(header):
    total = 6.5
    if len(header) == 2:
        return [2.0, 4.5]
    share = total / len(header)
    return [share] * len(header)


def add_heading(doc, level, text):
    heading = doc.add_heading(level=min(level, 3))
    run = heading.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(HEADING_SIZES.get(level, 11))
    run.font.bold = True
    run.font.color.rgb = NAVY if level <= 2 else BLUE


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    run = cell.paragraphs[0].add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(241, 245, 249)
    doc.add_paragraph()


def add_list_item(doc, text, ordered):
    paragraph = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    add_inline(paragraph, text, size=9.5)


def add_title_block(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("MediLocker")
    run.font.name = "Segoe UI"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Implementation Plan — Completion of Outstanding Scope\n")
    run.font.name = "Segoe UI"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = BLUE

    run = subtitle.add_run(
        "Six-phase plan covering the remaining 40% of the specified system, "
        "from defect repair to secure deployment"
    )
    run.font.name = "Segoe UI"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = GREY

    add_styled_table(
        doc,
        ["Document Property", "Detail"],
        [
            ["Project", "MediLocker — Secure Cloud-Based Digital Medical Record Management Platform"],
            ["Document", "Implementation Plan — Completion of Outstanding Scope"],
            ["Version & Status", "Version 2.0 — Draft for approval"],
            ["Date", "13 September 2026"],
            ["Scope", "Remaining work only (~40%); delivered scope is not restated"],
            ["Phases", "6 phases, 58 tasks"],
            ["Critical Path", "Phase 1 → 2 → 3 → 5 (4–6 weeks, single developer)"],
            ["Companion Documents",
             "DATABASE_SCHEMA_AND_ARCHITECTURE.md, ROLE_AND_PAGE_WISE_FEATURE_MAP.md"],
        ],
        col_widths=[2.0, 4.5],
    )


# ---------------------------------------------------------------------- build


def build():
    with open(SOURCE, "r", encoding="utf-8") as handle:
        lines = handle.read().splitlines()

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    add_title_block(doc)

    index = 0
    seen_title = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, "\n".join(block))
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if level == 1 and not seen_title:
                seen_title = True  # the title block above already covers it
                index += 1
                continue
            add_heading(doc, level, text)
            index += 1
            continue

        if stripped.startswith("|"):
            header, rows, index = parse_table(lines, index)
            if rows:
                add_styled_table(doc, header, rows, col_widths=column_widths(header))
            continue

        ordered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if ordered:
            add_list_item(doc, ordered.group(1), ordered=True)
            index += 1
            continue

        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:], ordered=False)
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, stripped)
        index += 1

    os.makedirs(os.path.join(ROOT, "docs"), exist_ok=True)
    for path in OUTPUTS:
        doc.save(path)
        print(f"Wrote {path}")


if __name__ == "__main__":
    build()
